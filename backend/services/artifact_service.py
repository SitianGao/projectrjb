"""将 Agent 生成的大纲/讲稿转换为真实 PPT 与音频文件。"""

from __future__ import annotations

import base64
import hashlib
import hmac
import json
import re
from email.utils import formatdate
from pathlib import Path
from urllib.parse import urlencode

from config import (
    PROJECT_ROOT,
    XFYUN_TTS_API_KEY,
    XFYUN_TTS_API_SECRET,
    XFYUN_TTS_APP_ID,
    XFYUN_TTS_VOICE,
)


GENERATED_RESOURCE_DIR = PROJECT_ROOT / "data" / "generated_resources"
GENERATED_RESOURCE_DIR.mkdir(parents=True, exist_ok=True)


class ArtifactService:
    """严格生成文件；生成失败时抛错，不伪装成成功的文本资源。"""

    def create(self, resource_id: str, resource_type: str, title: str, content: str):
        if resource_type == "ppt":
            path = self._create_ppt(resource_id, title, content)
            return f"/generated/{path.name}", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        if resource_type == "audio":
            path = self._create_audio(resource_id, content)
            return f"/generated/{path.name}", "audio/mpeg"
        return None, None

    @staticmethod
    def _create_ppt(resource_id: str, title: str, content: str) -> Path:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("生成 PPT 需要安装 python-pptx") from exc

        presentation = Presentation()
        cover = presentation.slides.add_slide(presentation.slide_layouts[0])
        cover.shapes.title.text = title
        cover.placeholders[1].text = "EduAgent · 由学习资源智能体生成"

        sections = re.split(r"(?=^\s*Slide\s+\d+\s*[:：])", content, flags=re.MULTILINE)
        sections = [section.strip() for section in sections if section.strip()]
        if not sections:
            sections = [content]
        for index, section in enumerate(sections[:12], 1):
            lines = [line.strip() for line in section.splitlines() if line.strip()]
            heading = re.sub(r"^Slide\s+\d+\s*[:：]\s*", "", lines[0]) if lines else f"第 {index} 页"
            bullets = [re.sub(r"^[-•]\s*", "", line) for line in lines[1:] if "讲师备注" not in line]
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = heading[:80]
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for bullet_index, bullet in enumerate((bullets or ["请结合学习资源正文完成本页学习"])[:6]):
                paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
                paragraph.text = bullet[:240]

        path = GENERATED_RESOURCE_DIR / f"{resource_id}.pptx"
        presentation.save(path)
        return path

    @staticmethod
    def _create_audio(resource_id: str, content: str) -> Path:
        if not all([XFYUN_TTS_APP_ID, XFYUN_TTS_API_KEY, XFYUN_TTS_API_SECRET]):
            raise RuntimeError(
                "audio 资源需要配置 XFYUN_TTS_APP_ID、XFYUN_TTS_API_KEY、XFYUN_TTS_API_SECRET"
            )
        try:
            import websocket
        except ImportError as exc:
            raise RuntimeError("生成音频需要安装 websocket-client") from exc

        host = "tts-api.xfyun.cn"
        request_path = "/v2/tts"
        date = formatdate(usegmt=True)
        signature_origin = f"host: {host}\ndate: {date}\nGET {request_path} HTTP/1.1"
        signature = base64.b64encode(
            hmac.new(
                XFYUN_TTS_API_SECRET.encode("utf-8"),
                signature_origin.encode("utf-8"),
                digestmod=hashlib.sha256,
            ).digest()
        ).decode("utf-8")
        authorization_origin = (
            f'api_key="{XFYUN_TTS_API_KEY}", algorithm="hmac-sha256", '
            f'headers="host date request-line", signature="{signature}"'
        )
        authorization = base64.b64encode(authorization_origin.encode("utf-8")).decode("utf-8")
        url = f"wss://{host}{request_path}?{urlencode({'authorization': authorization, 'date': date, 'host': host})}"

        # 官方接口单次文本限制小于 8000 字节，按 UTF-8 字节安全截取。
        text_bytes = content.encode("utf-8")[:7600]
        narration = text_bytes.decode("utf-8", errors="ignore")
        request = {
            "common": {"app_id": XFYUN_TTS_APP_ID},
            "business": {
                "aue": "lame",
                "auf": "audio/L16;rate=16000",
                "vcn": XFYUN_TTS_VOICE,
                "tte": "utf8",
            },
            "data": {
                "status": 2,
                "text": base64.b64encode(narration.encode("utf-8")).decode("utf-8"),
            },
        }
        connection = websocket.create_connection(url, timeout=30)
        chunks = []
        try:
            connection.send(json.dumps(request, ensure_ascii=False))
            while True:
                response = json.loads(connection.recv())
                if response.get("code") != 0:
                    raise RuntimeError(f"讯飞语音合成失败: {response.get('message') or response.get('code')}")
                audio = (response.get("data") or {}).get("audio")
                if audio:
                    chunks.append(base64.b64decode(audio))
                if (response.get("data") or {}).get("status") == 2:
                    break
        finally:
            connection.close()
        if not chunks:
            raise RuntimeError("讯飞语音合成未返回音频数据")
        path = GENERATED_RESOURCE_DIR / f"{resource_id}.mp3"
        path.write_bytes(b"".join(chunks))
        return path


artifact_service = ArtifactService()
